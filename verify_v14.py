"""Scripted verification for CUVAS CMS v0.14 — Step 19 reports,
audit log & backup:
  * the Reports centre (admin-only) with four reports whose figures
    are checked against INDEPENDENTLY computed ledger / voucher /
    results / attendance sums over a controlled fixture (three
    students, one fee plan of Rs 15,000: one pays in full, one only
    submits a proof, one splits and pays one installment) — issued
    Rs 45,000, collected Rs 22,500, outstanding Rs 22,500, the
    session-students cross-check equal, defaulters ordered B then C;
  * student strength counts equal SQL GROUP BY counts (gender,
    session, program x session matrix);
  * results summary: marks 80 / 40 / absent -> appeared 3, passed 1,
    failed 2, pass % 33.3, average 60.0, highest 80, Published Yes;
  * attendance summary equals the shared attendance_summary() output
    (2 classes, average 50.0%);
  * every Excel export downloads and opens with the expected sheets;
  * the append-only audit log records a scripted sequence (failed
    login, login, student create, voucher generate + verify, marks
    publish, campus settings save, student ticket, document upload,
    backup download, logout) with the right username + action, and
    its filters (text / role / date) narrow correctly;
  * the backup ZIP is valid, its workbook has no password_hash in
    the users sheet, stored blobs travel under files/, the README
    states the no-passwords + manual-restore notes, and a
    backup_downloaded audit row exists afterwards;
  * empty-session reports render cleanly; the student dashboard
    keeps exactly its 10 byte-identical tabs.

Expects a FRESH database (empty users table) so seeding + counts are exact.
Env: DB_PATH=<fresh sqlite file>  OR  DATABASE_URL=<postgres uri>.
Prints PASS/FAIL lines and exits non-zero if any check fails.
"""
import datetime
import io
import sys
import zipfile

from openpyxl import load_workbook

import app as m

FAIL = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail else ""))
    if not cond:
        FAIL.append(name)


def login(client, username, password):
    return client.post("/login", data={"username": username, "password": password})


def dbq(sql, params=()):
    with m.app.app_context():
        return m.get_db().execute(sql, params).fetchall()


def db1(sql, params=()):
    rows = dbq(sql, params)
    return rows[0] if rows else None


def audit_rows(action=None, username=None):
    sql, params = "SELECT * FROM audit_log WHERE 1=1", []
    if action:
        sql += " AND action=?"
        params.append(action)
    if username:
        sql += " AND username=?"
        params.append(username)
    return dbq(sql + " ORDER BY id", params)


def xlsx(resp):
    return load_workbook(io.BytesIO(resp.data))


def rows_of(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


NEW_URLS = ["/admin/reports", "/admin/reports/fees",
            "/admin/reports/fees.xlsx", "/admin/reports/students",
            "/admin/reports/students.xlsx", "/admin/reports/results",
            "/admin/reports/results.xlsx", "/admin/reports/attendance",
            "/admin/reports/attendance.xlsx", "/admin/audit",
            "/admin/backup", "/admin/backup/download"]

# ------------------------------------------------------- guards
anon = m.app.test_client()
check("anon: all v0.14 pages redirect to login",
      all(anon.get(u).status_code == 302 for u in NEW_URLS))
stu0 = m.app.test_client()
login(stu0, "student", "Student@123")
check("student: all v0.14 pages are refused (redirect)",
      all(stu0.get(u).status_code == 302 for u in NEW_URLS))
fac0 = m.app.test_client()
login(fac0, "faculty", "Faculty@123")
check("faculty: all v0.14 pages are refused (redirect)",
      all(fac0.get(u).status_code == 302 for u in NEW_URLS))

# ------------------------------------------------------- audit: login events
bad = m.app.test_client()
bad.post("/login", data={"username": "admin", "password": "WRONG"})
fails = audit_rows("login_failed", "admin")
check("failed login is audited with the attempted username",
      len(fails) == 1 and fails[0]["role"] == "admin"
      and fails[0]["user_id"] is not None)
admin = m.app.test_client()
r = login(admin, "admin", "Admin@123")
check("admin login works", r.status_code == 302)
oks = audit_rows("login_success", "admin")
check("successful login is audited (username + role snapshot)",
      len(oks) == 1 and oks[0]["role"] == "admin")

hub = admin.get("/admin/reports").get_data(as_text=True)
check("reports hub renders with the four reports + audit + backup cards",
      all(t in hub for t in ("Fee Collection", "Student Strength",
                             "Results Summary", "Attendance Summary",
                             "Audit Log", "Backup"))
      and "Generated on" in hub)
check("admin sidebar links the Reports hub",
      "/admin/reports" in admin.get("/admin").get_data(as_text=True))

# ------------------------------------------------------- fixtures
admin.post("/admin/departments/new", data={
    "name": "Dept of Reports (V14)", "code": "RPT-V14",
    "faculty_name": "Faculty (V14)"})
dept = db1("SELECT * FROM departments WHERE code='RPT-V14'")
admin.post("/admin/programs/new", data={
    "name": "BS Reports (V14)", "department_id": dept["id"],
    "degree_level": "BS", "duration_semesters": 8})
prog = db1("SELECT * FROM programs WHERE name='BS Reports (V14)'")
admin.post("/admin/sessions/new", data={
    "name": "Fall 2099 (V14)", "starts_on": "2099-09-01",
    "ends_on": "2100-01-15", "is_open": "1"})
sess = db1("SELECT * FROM academic_sessions WHERE name='Fall 2099 (V14)'")
admin.post("/admin/courses/new", data={
    "program_id": prog["id"], "semester_no": 1, "code": "RPT-101",
    "title": "Reporting Basics (V14)", "credit_theory": 3,
    "credit_lab": 0, "faculty_name": ""})
course = db1("SELECT * FROM courses WHERE code='RPT-101'")
check("fixtures created (dept / program / session / course)",
      all(x is not None for x in (dept, prog, sess, course)))

FIX = [("V14-REG-001", "V14 Student A", "Male", "0300-1111111"),
       ("V14-REG-002", "V14 Student B", "Female", "0300-2222222"),
       ("V14-REG-003", "V14 Student C", "Male", "")]
students = {}
for reg, name, gender, phone in FIX:
    admin.post("/admin/students/new", data={
        "registration_no": reg, "name": name, "gender": gender,
        "phone": phone, "department_id": dept["id"],
        "program_id": prog["id"], "session_id": sess["id"],
        "status": "Active", "current_semester": 1,
        "initial_password": "Test@1234"})
    students[reg] = db1("SELECT * FROM students WHERE registration_no=?",
                        (reg,))
check("fixture students A/B/C created",
      all(students[r] is not None for r, *_ in FIX))
created = audit_rows("student_created", "admin")
check("student creates are audited (3, by admin)",
      len([r for r in created if "V14-REG" in (r["details"] or "")]) == 3,
      str(len(created)))
A, B, C = (students["V14-REG-001"], students["V14-REG-002"],
           students["V14-REG-003"])

# ------------------------------------------------------- fees fixture
heads = dbq("SELECT * FROM fee_heads ORDER BY id")
admin.post("/admin/fee-plans/new", data={
    "program_id": prog["id"], "session_id": sess["id"], "semester_no": 1,
    "title": "V14 Plan", "due_date": "2099-11-01",
    "line_head_id": [str(heads[0]["id"]), str(heads[1]["id"])],
    "line_amount": ["10000", "5000"]})
plan = db1("SELECT * FROM fee_plans WHERE title='V14 Plan'")
check("fee plan created with total Rs 15,000",
      plan is not None and db1(
          "SELECT COALESCE(SUM(amount),0) t FROM fee_plan_lines "
          "WHERE plan_id=?", (plan["id"],))["t"] == 15000)
admin.post("/admin/vouchers/generate", data={
    "plan_id": plan["id"], "due_date": "2099-11-01", "student_id": "all"})
vouch = {s["id"]: db1("SELECT * FROM vouchers WHERE fee_plan_id=? AND "
                      "student_id=?", (plan["id"], s["id"]))
         for s in (A, B, C)}
check("3 vouchers generated for the fixture session",
      all(vouch[s["id"]] is not None for s in (A, B, C)))
check("voucher generation is audited",
      len(audit_rows("voucher_generated", "admin")) >= 1)


def client_for(reg):
    c = m.app.test_client()
    login(c, reg, "Test@1234")
    return c


def upload_proof(client, vid, name):
    return client.post(
        f"/student/vouchers/{vid}/proof",
        data={"proof": (io.BytesIO(b"proof-bytes-" + name.encode()), name),
              "payment_channel": "Bank"})


stuA, stuB, stuC = (client_for("V14-REG-001"), client_for("V14-REG-002"),
                    client_for("V14-REG-003"))
upload_proof(stuA, vouch[A["id"]]["id"], "a-proof.png")
admin.post(f"/admin/vouchers/{vouch[A['id']]['id']}/verify")
check("A's voucher verified Paid",
      db1("SELECT status FROM vouchers WHERE id=?",
          (vouch[A["id"]]["id"],))["status"] == "Paid")
check("voucher verification is audited",
      len(audit_rows("voucher_verified", "admin")) == 1)

upload_proof(stuB, vouch[B["id"]]["id"], "b-proof.png")  # never verified
stuC.post(f"/student/vouchers/{vouch[C['id']]['id']}/installments")
children = dbq("SELECT * FROM vouchers WHERE parent_voucher_id=? ORDER BY id",
               (vouch[C["id"]]["id"],))
check("C's voucher split into 2 x Rs 7,500",
      len(children) == 2
      and all(ch["amount"] == 7500 for ch in children))
upload_proof(stuC, children[0]["id"], "c-proof.png")
admin.post(f"/admin/vouchers/{children[0]['id']}/verify")
check("C's first installment verified Paid",
      db1("SELECT status FROM vouchers WHERE id=?",
          (children[0]["id"],))["status"] == "Paid")

# --- independent recomputation straight from the ledger
ind = db1("""SELECT COALESCE(SUM(l.debit),0) d, COALESCE(SUM(l.credit),0) c
             FROM ledger_entries l JOIN vouchers v ON v.id=l.voucher_id
             WHERE v.session_id=?""", (sess["id"],))
check("independent ledger sums: issued 45,000 / collected 22,500",
      ind["d"] == 45000 and ind["c"] == 22500, f"{ind['d']}/{ind['c']}")

# --- fee report: page
page = admin.get("/admin/reports/fees",
                 query_string={"session_id": sess["id"]}).get_data(as_text=True)
check("fee report page shows issued 45,000 + collected 22,500 + outstanding 22,500",
      "45,000" in page and page.count("22,500") >= 2)
check("fee report defaulters: B before C, with phone + amounts",
      "V14 Student B" in page and "V14 Student C" in page
      and page.index("V14 Student B") < page.index("V14 Student C")
      and "0300-2222222" in page and "15,000" in page and "7,500" in page)
check("fee report page generated-on timestamp",
      "Generated on" in page)

# --- fee report: Excel
wb = xlsx(admin.get("/admin/reports/fees.xlsx",
                    query_string={"session_id": sess["id"]}))
check("fee xlsx sheets are Summary / By Department / Defaulters",
      wb.sheetnames == ["Summary", "By Department", "Defaulters"],
      str(wb.sheetnames))
summary = {r[0]: r[1] for r in rows_of(wb["Summary"])
           if r and isinstance(r[0], str) and r[1] is not None}
check("fee xlsx summary equals the ledger figures",
      summary.get("Challans Issued") == 4
      and summary.get("Total Issued (Rs)") == 45000
      and summary.get("Total Collected (Rs)") == 22500
      and summary.get("Total Outstanding (Rs)") == 22500
      and summary.get("Pending Proofs") == 1, str(summary))
dept_rows = rows_of(wb["By Department"])
fx = [r for r in dept_rows if r[0] == "Dept of Reports (V14)"]
check("fee xlsx department breakdown: 45,000 / 22,500 / 22,500",
      len(fx) == 1 and fx[0][1:] == [45000, 22500, 22500], str(fx))
def_rows = [r for r in rows_of(wb["Defaulters"])[1:] if r[0]]
check("fee xlsx defaulters ordered B (15,000) then C (7,500)",
      [r[1] for r in def_rows] == ["V14-REG-002", "V14-REG-003"]
      and [r[4] for r in def_rows] == [15000, 7500], str(def_rows))

# ------------------------------------------------------- students report
total_sql = db1("SELECT COUNT(*) c FROM students")["c"]
check("fixture total is 9 students (6 seeded + 3)", total_sql == 9)
wb = xlsx(admin.get("/admin/reports/students.xlsx"))
check("students xlsx sheets",
      wb.sheetnames == ["Summary", "By Department", "By Program",
                        "By Session", "By Gender", "Program x Session"],
      str(wb.sheetnames))
summ = {r[0]: r[1] for r in rows_of(wb["Summary"])
        if r and isinstance(r[0], str) and r[1] is not None}
check("students xlsx total equals SQL count",
      summ.get("Total Students") == total_sql)
gender_sql = {r["g"]: r["c"] for r in dbq(
    "SELECT COALESCE(NULLIF(gender,''),'Not recorded') g, COUNT(*) c "
    "FROM students GROUP BY g")}
gender_x = {r[0]: r[1] for r in rows_of(wb["By Gender"])[1:] if r[0]}
check("students xlsx gender counts equal SQL (Male 5 / Female 4)",
      gender_x == gender_sql and gender_x.get("Male") == 5
      and gender_x.get("Female") == 4, str(gender_x))
sess_x = {r[0]: r[1] for r in rows_of(wb["By Session"])[1:] if r[0]}
check("students xlsx session count: Fall 2099 (V14) = 3",
      sess_x.get("Fall 2099 (V14)") == 3)
mx = rows_of(wb["Program x Session"])
hdr = mx[0]
col = hdr.index("Fall 2099 (V14)")
row = [r for r in mx[1:] if r[0] == "BS Reports (V14)"]
check("students xlsx program x session matrix cell = 3",
      len(row) == 1 and row[0][col] == 3 and row[0][-1] == 3)
spage = admin.get("/admin/reports/students").get_data(as_text=True)
check("students report page shows the total + breakdowns",
      "Total Students: 9" in spage and "By Gender" in spage
      and "Program × Session" in spage)

# ------------------------------------------------------- results fixture + report
r = admin.post("/admin/results/sheet",
               query_string={"course_id": course["id"],
                             "session_id": sess["id"]},
               data={f"marks_{A['id']}": "80", f"marks_{B['id']}": "40",
                     f"absent_{C['id']}": "1"})
check("marks sheet save redirects", r.status_code == 302)
check("marks save is audited (results_saved)",
      len(audit_rows("results_saved", "admin")) >= 1)
admin.post("/admin/results/sheet/publish",
           data={"course_id": course["id"], "session_id": sess["id"],
                 "action": "publish"})
check("marks publish is audited (results_published)",
      len(audit_rows("results_published", "admin")) == 1)
wb = xlsx(admin.get("/admin/reports/results.xlsx",
                    query_string={"session_id": sess["id"]}))
check("results xlsx sheet name", wb.sheetnames == ["Results Summary"])
rrow = [r for r in rows_of(wb["Results Summary"]) if r and r[0] == "RPT-101"]
check("results row: appeared 3, passed 1, failed 2, 33.3%, avg 60, high 80, Yes",
      len(rrow) == 1 and rrow[0][3] == 3 and rrow[0][4] == 1
      and rrow[0][5] == 2 and rrow[0][6] == 33.3
      and rrow[0][7] == 60 and rrow[0][8] == 80
      and rrow[0][9] == "Yes", str(rrow))
rpage = admin.get("/admin/reports/results",
                  query_string={"session_id": sess["id"]}).get_data(as_text=True)
check("results report page shows the course + pass %",
      "RPT-101" in rpage and "33.3%" in rpage)

# ------------------------------------------------------- attendance fixture + report
teacher = db1("SELECT * FROM faculty WHERE email='demo.teacher@cuvas.edu.pk'")
admin.post(f"/admin/faculty/{teacher['id']}/assign",
           data={"course_id": course["id"], "session_id": sess["id"]})
asg = db1("SELECT * FROM course_assignments WHERE faculty_id=? AND "
          "course_id=? AND session_id=?",
          (teacher["id"], course["id"], sess["id"]))
check("fixture course assigned to the demo teacher", asg is not None)
fac = m.app.test_client()
login(fac, "faculty", "Faculty@123")
today = datetime.date.today()
d1 = (today - datetime.timedelta(days=3)).isoformat()
d2 = (today - datetime.timedelta(days=2)).isoformat()
fac.post(f"/faculty/course/{asg['id']}/attendance/mark",
         query_string={"date": d1},
         data={f"status_{A['id']}": "Present",
               f"status_{B['id']}": "Present",
               f"status_{C['id']}": "Absent"})
fac.post(f"/faculty/course/{asg['id']}/attendance/mark",
         query_string={"date": d2},
         data={f"status_{A['id']}": "Present",
               f"status_{B['id']}": "Absent",
               f"status_{C['id']}": "Absent"})
with m.app.app_context():
    summ_att = m.attendance_summary(m.get_db(), asg["id"])
pcts = sorted(summ_att["by_student"][s["id"]]["pct"] for s in (A, B, C))
check("shared attendance_summary: 2 classes, pcts 0 / 50 / 100",
      summ_att["classes_held"] == 2 and pcts == [0.0, 50.0, 100.0],
      str(pcts))
wb = xlsx(admin.get("/admin/reports/attendance.xlsx",
                    query_string={"session_id": sess["id"]}))
check("attendance xlsx sheet name",
      wb.sheetnames == ["Attendance Summary"])
arow = [r for r in rows_of(wb["Attendance Summary"])
        if r and r[0] == "RPT-101"]
check("attendance row equals attendance_summary (2 held, 3 students, avg 50.0)",
      len(arow) == 1 and arow[0][4] == summ_att["classes_held"]
      and arow[0][5] == 3 and arow[0][6] == 50.0, str(arow))

# ------------------------------------------------------- student-side audit events
stuA.post("/student/tickets/new", data={
    "category": "Fee", "subject": "V14 ticket", "message": "Please help"})
tk = audit_rows("ticket_created", "V14-REG-001")
check("student ticket creation is audited with the student as actor",
      len(tk) == 1 and tk[0]["role"] == "student")
stuA.post("/student/documents/upload", data={
    "doc_type": "CNIC Front",
    "file": (io.BytesIO(b"\x89PNG-v14-doc-bytes"), "v14-cnic.png",
             "image/png")})
check("student document upload is audited",
      len(audit_rows("document_uploaded", "V14-REG-001")) == 1)

# ------------------------------------------------------- settings save audit
admin.post("/admin/campus", data={
    "lms_url": "https://lms.example.edu.pk", "library_url": "",
    "library_info": "", "email_url": "", "email_info": ""})
check("campus settings save is audited",
      len(audit_rows("campus_settings_saved", "admin")) == 1)

# ------------------------------------------------------- backup
r = admin.get("/admin/backup/download")
check("backup download is a ZIP attachment named cuvas-backup-*.zip",
      r.status_code == 200 and "application/zip" in r.headers["Content-Type"]
      and "cuvas-backup-" in r.headers.get("Content-Disposition", "")
      and ".zip" in r.headers.get("Content-Disposition", ""))
zf = zipfile.ZipFile(io.BytesIO(r.data))
names = zf.namelist()
check("backup ZIP contains the workbook + README.txt",
      "cuvas-backup.xlsx" in names and "README.txt" in names)
check("backup ZIP carries stored blobs under files/ (doc + proofs)",
      any(n.startswith("files/student_documents-") for n in names)
      and any(n.startswith("files/vouchers-") for n in names),
      str([n for n in names if n.startswith("files/")]))
bwb = load_workbook(io.BytesIO(zf.read("cuvas-backup.xlsx")))
check("backup workbook has the core sheets",
      all(s in bwb.sheetnames for s in
          ("departments", "programs", "sessions", "courses", "students",
           "faculty", "users", "enrollments", "vouchers", "ledger",
           "results", "attendance", "tickets", "student_documents",
           "audit_log")), str(bwb.sheetnames))
uhead = [c for c in rows_of(bwb["users"])[0]]
check("backup users sheet has NO password_hash column",
      "password_hash" not in uhead and "username" in uhead, str(uhead))
check("backup students sheet holds every student (+ header)",
      bwb["students"].max_row == total_sql + 1)
readme = zf.read("README.txt").decode("utf-8")
check("backup README states no-passwords + manual restore",
      "Passwords are NEVER included" in readme
      and "Restore is manual" in readme and "audit_log" in readme)
check("backup download is audited (backup_downloaded)",
      len(audit_rows("backup_downloaded", "admin")) == 1)
bpage = admin.get("/admin/backup").get_data(as_text=True)
check("backup page shows the last-backup timestamp afterwards",
      "Last backup downloaded" in bpage
      and datetime.date.today().isoformat() in bpage)

# ------------------------------------------------------- audit viewer + filters
apage = admin.get("/admin/audit").get_data(as_text=True)
check("audit page renders the recorded actions",
      all(a in apage for a in ("login_success", "student_created",
                               "voucher_verified", "backup_downloaded")))
f1 = admin.get("/admin/audit", query_string={"q": "student_created"}
               ).get_data(as_text=True)
check("audit text filter narrows (student_created only)",
      "student_created" in f1 and "voucher_verified" not in f1
      and "backup_downloaded" not in f1)
f2 = admin.get("/admin/audit", query_string={"role": "student"}
               ).get_data(as_text=True)
check("audit role filter narrows (student actions only)",
      "ticket_created" in f2 and "student_created" not in f2)
f3 = admin.get("/admin/audit", query_string={"date_from": "2999-01-01"}
               ).get_data(as_text=True)
check("audit future date range shows the clean empty state",
      "No audit entries match" in f3)
f4 = admin.get("/admin/audit",
               query_string={"date_from": datetime.date.today().isoformat()}
               ).get_data(as_text=True)
check("audit date_from=today still finds today's rows",
      "student_created" in f4)

# ------------------------------------------------------- logout audit
solo = m.app.test_client()
login(solo, "admin", "Admin@123")
solo.get("/logout")
check("logout is audited",
      len(audit_rows("logout", "admin")) == 1)

# ------------------------------------------------------- empty-session reports
admin.post("/admin/sessions/new", data={
    "name": "Empty 2099 (V14)", "starts_on": "", "ends_on": ""})
empty = db1("SELECT * FROM academic_sessions WHERE name='Empty 2099 (V14)'")
fe = admin.get("/admin/reports/fees", query_string={"session_id": empty["id"]})
check("empty session: fee report renders cleanly with empty states",
      fe.status_code == 200
      and "No challans were issued in this session" in fe.get_data(as_text=True)
      and "No defaulters" in fe.get_data(as_text=True))
re_ = admin.get("/admin/reports/results",
                query_string={"session_id": empty["id"]})
check("empty session: results report renders cleanly",
      re_.status_code == 200
      and "No results were recorded in this session"
      in re_.get_data(as_text=True))
at = admin.get("/admin/reports/attendance",
               query_string={"session_id": empty["id"]})
check("empty session: attendance report renders cleanly",
      at.status_code == 200
      and "No course assignments exist in this session"
      in at.get_data(as_text=True))
for u in ("/admin/reports/fees.xlsx", "/admin/reports/results.xlsx",
          "/admin/reports/attendance.xlsx"):
    r = admin.get(u, query_string={"session_id": empty["id"]})
    check(f"empty session: {u} still downloads a valid workbook",
          r.status_code == 200 and load_workbook(io.BytesIO(r.data)))

# ------------------------------------------------------- student tabs intact
dash = stuA.get("/student").get_data(as_text=True)
LABELS = ["News Feed", "Profile", "Course Enrollments", "Plan Of Study",
          "Transcript", "Financial Voucher", "Fee Plan &amp; Scholarships",
          "Account Statement", "My Documents", "Support Tickets"]
check("student dashboard keeps exactly 10 tabs, labels byte-identical",
      dash.count("data-tab=") == 10
      and all(f">{label}<" in dash for label in LABELS))

print()
if FAIL:
    print(f"{len(FAIL)} FAILURES: {FAIL}")
    sys.exit(1)
print("ALL CHECKS PASSED")
