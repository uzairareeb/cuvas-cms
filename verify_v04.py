"""Scripted verification for CUVAS CMS v0.4 — results, transcripts, slips.

Expects a FRESH database (empty users table) so seeding + counts are exact.
Env: DB_PATH=<fresh sqlite file>  OR  DATABASE_URL=<postgres uri>.
Prints PASS/FAIL lines and exits non-zero if any check fails.
"""
import base64
import io
import re
import sys
import zlib

import app as m

FAIL = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail else ""))
    if not cond:
        FAIL.append(name)


def login(client, username, password):
    return client.post("/login", data={"username": username, "password": password})


def dbq(sql, params=()):
    with m.app.app_context():
        return m.get_db().execute(sql, params).fetchall()


def pdf_text(pdf_bytes):
    """Decode ReportLab content streams (ASCII85 + Flate) to searchable text."""
    blob = b""
    for s in re.findall(rb"stream\r?\n(.*?)endstream", pdf_bytes, re.S):
        try:
            blob += zlib.decompress(base64.a85decode(s.strip(), adobe=True))
        except Exception:
            try:
                blob += zlib.decompress(s.strip())
            except Exception:
                pass
    return blob


def pdf_pages(pdf_bytes):
    return len(re.findall(rb"/Type\s*/Page(?![s\w])", pdf_bytes))


def transcript_pane(html):
    part = html.split('id="pane-transcript"', 1)
    return part[1].split("<section", 1)[0] if len(part) > 1 else ""


admin = m.app.test_client()
check("landing 200", admin.get("/").status_code == 200)
check("anon /admin/results redirects", admin.get("/admin/results").status_code == 302)
check("anon /admin/grade-scale redirects", admin.get("/admin/grade-scale").status_code == 302)
check("anon /student/transcript.pdf redirects", admin.get("/student/transcript.pdf").status_code == 302)
r = login(admin, "admin", "Admin@123")
check("admin login", r.status_code == 302 and "/admin" in r.headers["Location"])

# ---------------------------------------------------------------- grade scale
scale_page = admin.get("/admin/grade-scale").get_data(as_text=True)
check("grade scale page lists seeded bands",
      all(f'value="{g}"' in scale_page for g in
          ["A", "A-", "B+", "B", "B-", "C+", "C", "D", "F"]))
with m.app.app_context():
    bands = m.get_grade_bands(m.get_db())
check("9 default bands seeded", len(bands) == 9, str(len(bands)))


def derive(marks, absent=False):
    return m.grade_for_marks(marks, absent, bands)


check("derive 85 -> A/4.00", derive(85) == ("A", 4.0), str(derive(85)))
check("derive 84 -> A-/3.70", derive(84) == ("A-", 3.7), str(derive(84)))
check("derive 77 -> B+/3.30", derive(77) == ("B+", 3.3), str(derive(77)))
check("derive 74 -> B/3.00", derive(74) == ("B", 3.0), str(derive(74)))
check("derive 50 -> D/1.00", derive(50) == ("D", 1.0), str(derive(50)))
check("derive 49 -> F/0.00", derive(49) == ("F", 0.0), str(derive(49)))
check("derive absent -> F/0.00", derive(None, True) == ("F", 0.0), str(derive(None, True)))

# ------------------------------------------------- test class: dept/prog/session
admin.post("/admin/departments/new", data={
    "name": "Dept of Exams (V4)", "code": "EX-V4", "faculty_name": "Faculty (V4)"})
admin.post("/admin/sessions/new", data={
    "name": "Spring 2027 (V4)", "starts_on": "2027-02-01", "ends_on": "2027-06-30", "is_open": "1"})
with m.app.app_context():
    db = m.get_db()
    dept4 = db.execute("SELECT id FROM departments WHERE code='EX-V4'").fetchone()["id"]
    sess4 = db.execute("SELECT id FROM academic_sessions WHERE name='Spring 2027 (V4)'").fetchone()["id"]
admin.post("/admin/programs/new", data={
    "name": "BS Exams (V4)", "department_id": dept4, "degree_level": "BS", "duration_semesters": 8})
prog4 = dbq("SELECT id FROM programs WHERE name='BS Exams (V4)'")[0]["id"]

courses = {}
for code, th in [("V4-101", 3), ("V4-102", 3), ("V4-103", 2)]:
    admin.post("/admin/courses/new", data={
        "program_id": prog4, "semester_no": 1, "code": code,
        "title": f"Course {code} (V4)", "credit_theory": th, "credit_lab": 0,
        "faculty_name": "Dr V4"})
    courses[code] = dbq("SELECT id FROM courses WHERE code=?", (code,))[0]["id"]
check("3 test courses created", len(courses) == 3)

students = {}
for reg, name in [("V4-REG-001", "V4 Student One"), ("V4-REG-002", "V4 Student Two")]:
    r = admin.post("/admin/students/new", data={
        "registration_no": reg, "name": name, "department_id": dept4,
        "program_id": prog4, "session_id": sess4, "status": "Active",
        "current_semester": 1, "initial_password": "Test@1234"})
    check(f"student {reg} created", r.status_code == 302)
    students[reg] = dbq("SELECT id FROM students WHERE registration_no=?", (reg,))[0]["id"]
s1, s2 = students["V4-REG-001"], students["V4-REG-002"]


def sheet_save(course_id, form):
    return admin.post("/admin/results/sheet",
                      query_string={"course_id": course_id, "session_id": sess4}, data=form)


def sheet_publish(course_id, action):
    return admin.post("/admin/results/sheet/publish",
                      query_string={"course_id": course_id, "session_id": sess4},
                      data={"action": action})


# ------------------------------------------------------------- marks entry flow
sheet_save(courses["V4-101"], {f"marks_{s1}": "85", f"marks_{s2}": "70"})
row = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?", (s1, courses["V4-101"]))[0]
check("sheet save stores marks + derived grade A/4.0",
      row["marks_obtained"] == 85 and row["grade"] == "A" and row["grade_point"] == 4.0,
      str(dict(row)))
sheet_save(courses["V4-101"], {f"marks_{s1}": "88"})
rows = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?", (s1, courses["V4-101"]))
check("re-save updates in place (no duplicate, marks 88)",
      len(rows) == 1 and rows[0]["marks_obtained"] == 88, str(len(rows)))
sheet_save(courses["V4-102"], {f"marks_{s1}": "74", f"absent_{s2}": "on"})
row2 = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?", (s2, courses["V4-102"]))[0]
check("absent stored as absent + F/0.0",
      row2["is_absent"] == 1 and row2["marks_obtained"] is None and row2["grade"] == "F"
      and row2["grade_point"] == 0.0, str(dict(row2)))
enr = dbq("SELECT * FROM enrollments WHERE student_id=? AND course_id=?", (s1, courses["V4-101"]))
check("enrollment auto-created on result save", len(enr) == 1 and enr[0]["semester_no"] == 1)

# unpublished results must be invisible to the student
c1 = m.app.test_client()
login(c1, "V4-REG-001", "Test@1234")
pane = transcript_pane(c1.get("/student").get_data(as_text=True))
check("unpublished results hidden from student transcript",
      "V4-101" not in pane and "Overall CGPA: <strong>0.0</strong>" in pane)

sheet_publish(courses["V4-101"], "publish")
sheet_publish(courses["V4-102"], "publish")
pane = transcript_pane(c1.get("/student").get_data(as_text=True))
# S1: V4-101 88 -> A 4.0 x 3 CH = 12 ; V4-102 74 -> B 3.0 x 3 CH = 9 ; SGPA = 21/6 = 3.5
check("published results visible; SGPA 3.5 hand-check", "SGPA: 3.5" in pane, pane[:200])
check("web CGPA = 3.5", "Overall CGPA: <strong>3.5</strong>" in pane)
check("marks shown on transcript (88)", ">88<" in pane)

c2 = m.app.test_client()
login(c2, "V4-REG-002", "Test@1234")
pane2 = transcript_pane(c2.get("/student").get_data(as_text=True))
# S2: V4-101 70 -> B 3.0 x 3 = 9 ; V4-102 Absent F x 3 = 0 ; SGPA = 9/6 = 1.5
check("absent shows as Absent + F on transcript", "Absent" in pane2 and ">F<" in pane2)
check("S2 SGPA 1.5 hand-check", "SGPA: 1.5" in pane2, pane2[:200])

sheet_publish(courses["V4-101"], "unpublish")
pane = transcript_pane(c1.get("/student").get_data(as_text=True))
check("unpublish hides results again", "V4-101" not in pane and "SGPA: 3.0" in pane)
sheet_publish(courses["V4-101"], "publish")

landing = admin.get("/admin/results").get_data(as_text=True)
check("results landing lists set + Published state",
      "V4-101" in landing and "Published" in landing)

# derivation at the low band + clear/delete
sheet_save(courses["V4-103"], {f"marks_{s1}": "49"})
row3 = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?", (s1, courses["V4-103"]))[0]
check("49 marks stored as F/0.00", row3["grade"] == "F" and row3["grade_point"] == 0.0)
sheet_save(courses["V4-103"], {f"clear_{s1}": "on"})
check("clear deletes the result",
      dbq("SELECT * FROM results WHERE student_id=? AND course_id=?", (s1, courses["V4-103"])) == [])

# ----------------------------------------------------------------- marks Excel
dl = admin.get("/admin/results/sheet.xlsx",
               query_string={"course_id": courses["V4-103"], "session_id": sess4})
check("marks sheet xlsx downloads", dl.status_code == 200 and len(dl.data) > 500)
from openpyxl import Workbook, load_workbook
wb = load_workbook(io.BytesIO(dl.data))
ws = wb.active
sheet_rows = list(ws.iter_rows(values_only=True))
check("marks sheet has header + 2 class rows", len(sheet_rows) == 3, str(len(sheet_rows)))
# fill: S1 -> 90 (A), S2 -> 105 (out of range), plus an unknown reg row
for r in sheet_rows[1:]:
    if r[0] == "V4-REG-001":
        pass
ws.append(["V4-REG-001", "V4 Student One", 90, ""])
ws.append(["V4-REG-002", "V4 Student Two", 105, ""])
ws.append(["NOPE-999", "Ghost", 50, ""])
# remove the two original blank rows (rows 2-3) so only filled rows import
ws.delete_rows(2, 2)
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)
r = admin.post("/admin/results/import",
               query_string={"course_id": courses["V4-103"], "session_id": sess4},
               data={"file": (buf, "marks.xlsx")},
               content_type="multipart/form-data")
body = r.get_data(as_text=True)
check("import result page 200", r.status_code == 200)
check("import: 1 imported", ">1<" in body and "V4-REG-001" in body)
check("import: unknown reg skipped with reason",
      "NOPE-999" in body and "unknown registration no" in body)
check("import: out-of-range marks skipped with reason", "outside 0" in body)
row3 = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?", (s1, courses["V4-103"]))[0]
check("imported marks saved + derived A", row3["marks_obtained"] == 90 and row3["grade"] == "A")
check("skipped S2 got NO result from import",
      dbq("SELECT * FROM results WHERE student_id=? AND course_id=?", (s2, courses["V4-103"])) == [])

# ------------------------------------------------------------------- PDFs
t_pdf = c1.get("/student/transcript.pdf")
check("student transcript PDF %PDF + size", t_pdf.data[:5] == b"%PDF-" and len(t_pdf.data) > 1000,
      str(len(t_pdf.data)))
blob = pdf_text(t_pdf.data)
check("PDF CGPA text matches web (3.50)", b"3.50" in blob)
with m.app.app_context():
    _t, helper_cgpa, _ch = m.transcript_data(m.get_db(), s1, published_only=True)
check("helper CGPA == web CGPA == 3.5", helper_cgpa == 3.5, str(helper_cgpa))

slip1 = c1.get("/student/slip.pdf")
check("student own slip PDF 1 page", slip1.data[:5] == b"%PDF-" and pdf_pages(slip1.data) == 1,
      str(pdf_pages(slip1.data)))
slip_courses_blob = pdf_text(slip1.data)
check("slip lists the semester courses", b"V4-101" in slip_courses_blob and b"V4-102" in slip_courses_blob)
bulk = admin.get("/admin/slips.pdf", query_string={
    "program_id": prog4, "session_id": sess4, "semester_no": 1})
check("bulk slips PDF pages == student count (2)",
      bulk.data[:5] == b"%PDF-" and pdf_pages(bulk.data) == 2, str(pdf_pages(bulk.data)))
adm_slip = admin.get(f"/admin/students/{s2}/slip.pdf")
check("admin single slip PDF ok", adm_slip.status_code == 200 and adm_slip.data[:5] == b"%PDF-")
adm_tr = admin.get(f"/admin/students/{s1}/transcript.pdf")
check("admin transcript PDF ok", adm_tr.status_code == 200 and adm_tr.data[:5] == b"%PDF-")

# ------------------------------------------------------------------- guards
check("student cannot open admin transcript PDF",
      c1.get(f"/admin/students/{s2}/transcript.pdf").status_code == 302)
check("student cannot open admin slip PDF",
      c1.get(f"/admin/students/{s2}/slip.pdf").status_code == 302)
check("student cannot open bulk slips",
      c1.get("/admin/slips.pdf", query_string={
          "program_id": prog4, "session_id": sess4, "semester_no": 1}).status_code == 302)
check("student cannot open marks sheet",
      c1.get("/admin/results/sheet", query_string={
          "course_id": courses["V4-101"], "session_id": sess4}).status_code == 302)
check("S2 transcript does not leak S1's 88 marks",
      ">88<" not in transcript_pane(c2.get("/student").get_data(as_text=True)))

# ------------------------------------------------------- demo transcript intact
demo = m.app.test_client()
login(demo, "student", "Student@123")
dh = demo.get("/student").get_data(as_text=True)
dpane = transcript_pane(dh)
# v0.6 added the My Documents + Support Tickets tabs: 10 tabs total
check("demo transcript: 10 tabs + CGPA 3.63 intact",
      dh.count("data-tab=") == 10 and "Overall CGPA: <strong>3.63</strong>" in dpane)
check("demo SGPA semesters 3.61 / 3.65 intact",
      "SGPA: 3.61" in dpane and "SGPA: 3.65" in dpane, dpane[:120])
demo_pdf = demo.get("/student/transcript.pdf")
check("demo transcript PDF ok + shows 3.63",
      demo_pdf.data[:5] == b"%PDF-" and b"3.63" in pdf_text(demo_pdf.data))

# ------------------------------------------------- grade scale edit round-trip
with m.app.app_context():
    bands_now = m.get_grade_bands(m.get_db())
form = {"band_id": [str(b["id"]) for b in bands_now]}
for b in bands_now:
    form[f"min_{b['id']}"] = str(int(b["min_marks"]))
    form[f"grade_{b['id']}"] = b["grade"]
    form[f"gp_{b['id']}"] = str(b["grade_point"])
top = bands_now[0]
form[f"min_{top['id']}"] = "86"
r = admin.post("/admin/grade-scale", data=form)
check("grade scale edit saves", r.status_code == 302)
with m.app.app_context():
    bands2 = m.get_grade_bands(m.get_db())
check("after edit, 85 no longer an A", m.grade_for_marks(85, False, bands2)[0] == "A-",
      str(m.grade_for_marks(85, False, bands2)))
form[f"min_{top['id']}"] = "85"
admin.post("/admin/grade-scale", data=form)
with m.app.app_context():
    bands3 = m.get_grade_bands(m.get_db())
check("scale restored: 85 -> A again", m.grade_for_marks(85, False, bands3) == ("A", 4.0))
r = admin.post("/admin/grade-scale/add",
               data={"min_marks": "95", "grade": "A+", "grade_point": "4.0"})
with m.app.app_context():
    bands4 = m.get_grade_bands(m.get_db())
check("band add works (96 -> A+)", m.grade_for_marks(96, False, bands4)[0] == "A+")
plus = dbq("SELECT id FROM grade_bands WHERE grade='A+'")[0]["id"]
admin.post(f"/admin/grade-scale/{plus}/delete")
with m.app.app_context():
    bands5 = m.get_grade_bands(m.get_db())
check("band delete works (back to 9 bands)", len(bands5) == 9, str(len(bands5)))

print()
if FAIL:
    print(f"{len(FAIL)} FAILURES: {FAIL}")
    sys.exit(1)
print("ALL CHECKS PASSED")
