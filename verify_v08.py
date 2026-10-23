"""Scripted verification for CUVAS CMS v0.8 — Step 12 attendance:
faculty mark one sheet per class date for their OWN assignments only
(upsert per student+date — re-saving edits in place, never duplicates),
future dates rejected with zero writes, delete-a-date, and ONE shared
summary computation (classes held = distinct marked dates; percentage =
present / classes held) asserted IDENTICAL on the faculty summary, the
read-only admin overview and the student's own /student/attendance page.
Cross-faculty 404s on every attendance route (GET + POST), cross-student
isolation via exact (date, status) pair sets, the no-attendance empty
state, the summary Excel export, and role guards.

Expects a FRESH database (empty users table) so seeding + counts are exact.
Env: DB_PATH=<fresh sqlite file>  OR  DATABASE_URL=<postgres uri>.
Prints PASS/FAIL lines and exits non-zero if any check fails.
"""
import datetime
import io
import re
import sys

import app as m

FAIL = []

TODAY = datetime.date.today()
D1 = (TODAY - datetime.timedelta(days=10)).isoformat()
D2 = (TODAY - datetime.timedelta(days=9)).isoformat()
FUTURE = (TODAY + datetime.timedelta(days=30)).isoformat()


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail else ""))
    if not cond:
        FAIL.append(name)


def login(client, username, password):
    return client.post("/login", data={"username": username, "password": password})


def dbq(sql, params=()):
    with m.app.app_context():
        return m.get_db().execute(sql, params).fetchall()


def att_rows(aid, date=None):
    if date:
        return dbq("SELECT * FROM attendance WHERE assignment_id=? AND class_date=? "
                   "ORDER BY student_id", (aid, date))
    return dbq("SELECT * FROM attendance WHERE assignment_id=? ORDER BY class_date, student_id",
               (aid,))


def pairs(html):
    """(date, status) pairs rendered in a student attendance detail table."""
    return set(re.findall(r"<td>(\d{4}-\d{2}-\d{2})</td><td><span class=\"pill[^\"]*\">(\w+)</span>",
                          html))


admin = m.app.test_client()
check("landing 200", admin.get("/").status_code == 200)
check("anon /student/attendance redirects", admin.get("/student/attendance").status_code == 302)
check("anon /admin/attendance redirects", admin.get("/admin/attendance").status_code == 302)
r = login(admin, "admin", "Admin@123")
check("admin login", r.status_code == 302 and "/admin" in r.headers["Location"])

# ------------------------------------------------------- fixtures (V8 namespace)
admin.post("/admin/departments/new", data={
    "name": "Dept of Attendance (V8)", "code": "ATT-V8", "faculty_name": "Faculty (V8)"})
admin.post("/admin/sessions/new", data={
    "name": "Spring 2028 (V8)", "starts_on": "2028-02-01", "ends_on": "2028-06-30",
    "is_open": "1"})
dept8 = dbq("SELECT id FROM departments WHERE code='ATT-V8'")[0]["id"]
sess8 = dbq("SELECT id FROM academic_sessions WHERE name='Spring 2028 (V8)'")[0]["id"]
admin.post("/admin/programs/new", data={
    "name": "BS Attendance (V8)", "department_id": dept8, "degree_level": "BS",
    "duration_semesters": 8})
prog8 = dbq("SELECT id FROM programs WHERE name='BS Attendance (V8)'")[0]["id"]
courses8 = {}
for code in ("V8-101", "V8-102"):
    admin.post("/admin/courses/new", data={
        "program_id": prog8, "semester_no": 1, "code": code,
        "title": f"Course {code} (V8)", "credit_theory": 3, "credit_lab": 0,
        "faculty_name": ""})
    courses8[code] = dbq("SELECT id FROM courses WHERE code=?", (code,))[0]["id"]
students8 = {}
for reg, name in [("V8-REG-001", "V8 Student One"), ("V8-REG-002", "V8 Student Two"),
                  ("V8-REG-003", "V8 Student Three")]:
    admin.post("/admin/students/new", data={
        "registration_no": reg, "name": name, "department_id": dept8,
        "program_id": prog8, "session_id": sess8, "status": "Active",
        "current_semester": 1, "initial_password": "Test@1234"})
    students8[reg] = dbq("SELECT id FROM students WHERE registration_no=?", (reg,))[0]["id"]
s1, s2, s3 = (students8["V8-REG-001"], students8["V8-REG-002"],
              students8["V8-REG-003"])

admin.post("/admin/faculty/new", data={
    "full_name": "V8 Teacher A", "email": "teacher.a@v8.example",
    "designation": "Lecturer", "department_id": dept8,
    "initial_password": "Teach@123"})
admin.post("/admin/faculty/new", data={
    "full_name": "V8 Teacher B", "email": "teacher.b@v8.example",
    "designation": "Assistant Professor", "department_id": dept8,
    "initial_password": "Teach@456"})
fac_a = dbq("SELECT * FROM faculty WHERE email='teacher.a@v8.example'")[0]
fac_b = dbq("SELECT * FROM faculty WHERE email='teacher.b@v8.example'")[0]
admin.post(f"/admin/faculty/{fac_a['id']}/assign",
           data={"course_id": courses8["V8-101"], "session_id": sess8})
admin.post(f"/admin/faculty/{fac_b['id']}/assign",
           data={"course_id": courses8["V8-102"], "session_id": sess8})
aid_a = dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
            (fac_a["id"],))[0]["id"]
aid_b = dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
            (fac_b["id"],))[0]["id"]

fa = m.app.test_client()
login(fa, "teacher.a@v8.example", "Teach@123")
fb = m.app.test_client()
login(fb, "teacher.b@v8.example", "Teach@456")

# enrollments for students One/Two in V8-101 via the faculty marks flow
# (save_result creates them, exactly like the real marks workflow)
fa.post(f"/faculty/course/{aid_a}/marks",
        data={f"marks_{s1}": "70", f"marks_{s2}": "65"})
enr = dbq("SELECT student_id FROM enrollments WHERE course_id=?",
          (courses8["V8-101"],))
check("marks flow created enrollments for One+Two in V8-101",
      {r["student_id"] for r in enr} == {s1, s2}, str([dict(r) for r in enr]))
# student Three gets an enrollment in V8-102 (teacher B's course, never marked)
with m.app.app_context():
    raw = m.connect()
    raw.insert("INSERT INTO enrollments (student_id,course_id,session_name,semester_no) "
               "VALUES (?,?,?,?)", (s3, courses8["V8-102"], "Spring 2028 (V8)", 1))
    raw.commit()
    raw.conn.close()

# ------------------------------------------------------- faculty landing + mark flow
land = fa.get(f"/faculty/course/{aid_a}/attendance").get_data(as_text=True)
check("attendance landing 200, date picker defaults to today, empty state",
      f'value="{TODAY.isoformat()}"' in land and "No attendance has been marked" in land)
check("bad date rejected to landing",
      fa.get(f"/faculty/course/{aid_a}/attendance/mark",
             query_string={"date": "not-a-date"}).status_code == 302)
sheet = fa.get(f"/faculty/course/{aid_a}/attendance/mark",
               query_string={"date": D1}).get_data(as_text=True)
check("mark sheet lists the whole cohort with status selects",
      all(x in sheet for x in ("V8-REG-001", "V8-REG-002", "V8-REG-003",
                               f"status_{s1}", f"status_{s2}", f"status_{s3}"))
      and ">Present<" in sheet and ">Leave<" in sheet)

r = fa.post(f"/faculty/course/{aid_a}/attendance/mark", query_string={"date": D1},
            data={f"status_{s1}": "Present", f"status_{s2}": "Absent",
                  f"status_{s3}": "Late"})
check("mark date1 redirects", r.status_code == 302)
rows = att_rows(aid_a, D1)
by = {x["student_id"]: x["status"] for x in rows}
check("date1 rows saved with the chosen statuses",
      len(rows) == 3 and by == {s1: "Present", s2: "Absent", s3: "Late"},
      str(by))
check("rows stamped with the marking teacher",
      all(x["marked_by_faculty_id"] == fac_a["id"] for x in rows))
summ = fa.get(f"/faculty/course/{aid_a}/attendance/summary").get_data(as_text=True)
check("summary after date1: classes held 1, One 100.0%, Two 0.0%",
      "Classes held: <strong>1</strong>" in summ and "100.0%" in summ
      and "0.0%" in summ)
sheet = fa.get(f"/faculty/course/{aid_a}/attendance/mark",
               query_string={"date": D1}).get_data(as_text=True)
check("re-opening date1 pre-selects the saved statuses",
      'value="Absent" selected' in sheet and 'value="Late" selected' in sheet)

# ------------------------------------------------------- re-save = update in place
r = fa.post(f"/faculty/course/{aid_a}/attendance/mark", query_string={"date": D1},
            data={f"status_{s1}": "Absent", f"status_{s2}": "Present",
                  f"status_{s3}": "Leave"})
rows = att_rows(aid_a, D1)
by = {x["student_id"]: x["status"] for x in rows}
check("re-saving date1 updates in place (still 3 rows, new statuses)",
      len(rows) == 3 and by == {s1: "Absent", s2: "Present", s3: "Leave"},
      str(by))
summ = fa.get(f"/faculty/course/{aid_a}/attendance/summary").get_data(as_text=True)
check("classes_held still 1 after the re-save",
      "Classes held: <strong>1</strong>" in summ)

# canonical date1 + date2 -> hand-computed percentages
fa.post(f"/faculty/course/{aid_a}/attendance/mark", query_string={"date": D1},
        data={f"status_{s1}": "Present", f"status_{s2}": "Absent",
              f"status_{s3}": "Present"})
fa.post(f"/faculty/course/{aid_a}/attendance/mark", query_string={"date": D2},
        data={f"status_{s1}": "Absent", f"status_{s2}": "Present",
              f"status_{s3}": "Present"})
check("two dates -> 6 rows total, no duplicates",
      len(att_rows(aid_a)) == 6)
with m.app.app_context():
    summary = m.attendance_summary(m.get_db(), aid_a)
check("helper: classes_held 2; One 1P/1A = 50.0; Two 1A/1P = 50.0; Three 2P = 100.0",
      summary["classes_held"] == 2
      and summary["by_student"][s1]["present"] == 1
      and summary["by_student"][s1]["absent"] == 1
      and summary["by_student"][s1]["pct"] == 50.0
      and summary["by_student"][s2]["pct"] == 50.0
      and summary["by_student"][s3]["present"] == 2
      and summary["by_student"][s3]["pct"] == 100.0,
      str({k: v for k, v in summary["by_student"].items()}))
summ = fa.get(f"/faculty/course/{aid_a}/attendance/summary").get_data(as_text=True)
check("faculty summary page shows classes held 2 + 50.0% + 100.0%",
      "Classes held: <strong>2</strong>" in summ and "50.0%" in summ
      and "100.0%" in summ)
land = fa.get(f"/faculty/course/{aid_a}/attendance").get_data(as_text=True)
check("landing lists both dates with per-date counts",
      D1 in land and D2 in land)

# ------------------------------------------------------- Excel export
xl = fa.get(f"/faculty/course/{aid_a}/attendance/summary.xlsx")
check("summary Excel downloads", xl.status_code == 200 and xl.data[:2] == b"PK")
from openpyxl import load_workbook
ws = load_workbook(io.BytesIO(xl.data)).active
xl_rows = [list(r) for r in ws.iter_rows(values_only=True)]
hdr = next(i for i, r in enumerate(xl_rows) if r and r[0] == "Reg No")
data = {r[0]: r for r in xl_rows[hdr + 1:] if r and r[0]}
check("Excel: classes-held line + header + One's row (1,1,0,0,2,50.0)",
      any(r and isinstance(r[0], str) and r[0] == "Classes held: 2"
          for r in xl_rows)
      and xl_rows[hdr] == ["Reg No", "Name", "Present", "Absent", "Late",
                           "Leave", "Total", "Attendance %"]
      and data["V8-REG-001"][2:] == [1, 1, 0, 0, 2, 50.0]
      and data["V8-REG-003"][2:] == [2, 0, 0, 0, 2, 100.0],
      str(xl_rows))

# ------------------------------------------------------- student pages (same numbers)
stu1 = m.app.test_client()
login(stu1, "V8-REG-001", "Test@1234")
page1 = stu1.get("/student/attendance").get_data(as_text=True)
check("student One: V8-101 + teacher A + classes held 2 + 50.0%",
      "V8-101" in page1 and "V8 Teacher A" in page1
      and "Classes held: <strong>2</strong>" in page1 and "50.0%" in page1)
check("student One detail pairs are EXACTLY his own",
      pairs(page1) == {(D1, "Present"), (D2, "Absent")}, str(pairs(page1)))
check("student One page leaks nothing of Two (no name, no mirrored pairs)",
      "V8 Student Two" not in page1
      and (D1, "Absent") not in pairs(page1)
      and (D2, "Present") not in pairs(page1))
stu2 = m.app.test_client()
login(stu2, "V8-REG-002", "Test@1234")
page2 = stu2.get("/student/attendance").get_data(as_text=True)
check("student Two detail pairs are EXACTLY his own (mirror of One)",
      pairs(page2) == {(D1, "Absent"), (D2, "Present")}, str(pairs(page2)))
stu3 = m.app.test_client()
login(stu3, "V8-REG-003", "Test@1234")
page3 = stu3.get("/student/attendance").get_data(as_text=True)
check("student Three (course assigned, never marked): empty state shown",
      "V8-102" in page3 and "V8 Teacher B" in page3
      and "No attendance recorded yet for this course." in page3
      and pairs(page3) == set())
dash = stu1.get("/student").get_data(as_text=True)
check("student dashboard: 10 tabs intact + My Attendance link in enrollments",
      dash.count("data-tab=") == 10 and "My Attendance" in dash
      and 'href="/student/attendance"' in dash)

# ------------------------------------------------------- admin overview (read-only, same numbers)
over = admin.get("/admin/attendance").get_data(as_text=True)
check("admin overview lists both assignments with teacher + classes held",
      "V8 Teacher A" in over and "V8-101" in over and "V8-102" in over
      and f"/admin/attendance/{aid_a}" in over)
det = admin.get(f"/admin/attendance/{aid_a}").get_data(as_text=True)
check("admin detail: same summary (classes held 2, 50.0%, One's counts)",
      "Classes held: <strong>2</strong>" in det and "50.0%" in det
      and "V8-REG-001" in det and "V8 Teacher A" in det)
check("admin detail is read-only (no mark/delete controls)",
      "attendance/mark" not in det and "<form" not in det)
check("admin detail of unknown assignment -> 404",
      admin.get("/admin/attendance/999999").status_code == 404)

# ------------------------------------------------------- future dates rejected, zero writes
before = len(att_rows(aid_a))
r = fa.post(f"/faculty/course/{aid_a}/attendance/mark",
            query_string={"date": FUTURE},
            data={f"status_{s1}": "Present", f"status_{s2}": "Present",
                  f"status_{s3}": "Present"})
land = fa.get(f"/faculty/course/{aid_a}/attendance").get_data(as_text=True)
check("future-date POST rejected (redirect + flash), zero rows written",
      r.status_code == 302 and "in the future" in land
      and att_rows(aid_a, FUTURE) == [] and len(att_rows(aid_a)) == before)
check("future-date GET sheet also rejected",
      fa.get(f"/faculty/course/{aid_a}/attendance/mark",
             query_string={"date": FUTURE}).status_code == 302
      and att_rows(aid_a, FUTURE) == [])

# ------------------------------------------------------- cross-faculty 404s
n_before = len(att_rows(aid_a))
check("B landing of A's assignment -> 404",
      fb.get(f"/faculty/course/{aid_a}/attendance").status_code == 404)
check("B mark GET of A's assignment -> 404",
      fb.get(f"/faculty/course/{aid_a}/attendance/mark",
             query_string={"date": D1}).status_code == 404)
r = fb.post(f"/faculty/course/{aid_a}/attendance/mark",
            query_string={"date": D1}, data={f"status_{s1}": "Absent"})
check("B mark POST of A's assignment -> 404, nothing written",
      r.status_code == 404 and len(att_rows(aid_a)) == n_before
      and {x["student_id"]: x["status"] for x in att_rows(aid_a, D1)}[s1] == "Present")
check("B summary of A's assignment -> 404",
      fb.get(f"/faculty/course/{aid_a}/attendance/summary").status_code == 404)
check("B summary xlsx of A's assignment -> 404",
      fb.get(f"/faculty/course/{aid_a}/attendance/summary.xlsx").status_code == 404)
r = fb.post(f"/faculty/course/{aid_a}/attendance/delete", data={"date": D2})
check("B delete of A's date -> 404, rows intact",
      r.status_code == 404 and len(att_rows(aid_a, D2)) == 3)

# ------------------------------------------------------- role guards
check("faculty blocked from /admin/attendance", fa.get("/admin/attendance").status_code == 302)
check("faculty blocked from /student/attendance", fa.get("/student/attendance").status_code == 302)
check("student blocked from faculty attendance",
      stu1.get(f"/faculty/course/{aid_a}/attendance").status_code == 302)
check("student blocked from /admin/attendance", stu1.get("/admin/attendance").status_code == 302)
check("admin blocked from faculty attendance",
      admin.get(f"/faculty/course/{aid_a}/attendance").status_code == 302)
check("admin blocked from /student/attendance", admin.get("/student/attendance").status_code == 302)

# ------------------------------------------------------- delete a date -> recompute
r = fa.post(f"/faculty/course/{aid_a}/attendance/delete", data={"date": D2})
check("delete date2 redirects", r.status_code == 302)
check("date2 rows gone, date1 rows intact",
      att_rows(aid_a, D2) == [] and len(att_rows(aid_a, D1)) == 3)
summ = fa.get(f"/faculty/course/{aid_a}/attendance/summary").get_data(as_text=True)
check("after delete: classes held back to 1, One recomputed to 100.0%",
      "Classes held: <strong>1</strong>" in summ and "100.0%" in summ)
page1 = stu1.get("/student/attendance").get_data(as_text=True)
check("student One page recomputes too (100.0%, only date1 pair)",
      "100.0%" in page1 and pairs(page1) == {(D1, "Present")})
det = admin.get(f"/admin/attendance/{aid_a}").get_data(as_text=True)
check("admin detail recomputes too (classes held 1)",
      "Classes held: <strong>1</strong>" in det)

# ------------------------------------------------------- demo seed sanity
demo = m.app.test_client()
login(demo, "student", "Student@123")
dpage = demo.get("/student/attendance")
check("demo student attendance page 200 with the empty state",
      dpage.status_code == 200
      and "No attendance recorded yet" in dpage.get_data(as_text=True))
demo_fac = m.app.test_client()
login(demo_fac, "faculty", "Faculty@123")
ddash = demo_fac.get("/faculty").get_data(as_text=True)
check("demo teacher dashboard links Attendance per course",
      ddash.count("Attendance</a>") >= 2)

print()
if FAIL:
    print(f"{len(FAIL)} FAILURES: {FAIL}")
    sys.exit(1)
print("ALL CHECKS PASSED")
