"""Scripted verification for CUVAS CMS v0.7 — Step 11 faculty portal:
faculty records + auto-created logins (email = username, shown-once
passwords), course assignments (one primary teacher per course+session),
the /faculty portal (own courses only, roster + Excel export, marks entry
into the shared results store), the publish lock (faculty can never
publish; a published set is read-only for faculty), role guards,
deactivation, delete guard, the demo teacher seed, and a v0.6-era
database migration simulation.

Expects a FRESH database (empty users table) so seeding + counts are exact.
Env: DB_PATH=<fresh sqlite file>  OR  DATABASE_URL=<postgres uri>.
Prints PASS/FAIL lines and exits non-zero if any check fails.
"""
import io
import re
import sys

import app as m

FAIL = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail else ""))
    if not cond:
        FAIL.append(name)


def login(client, username, password):
    return client.post("/login", data={"username": username, "password": password})


def dbq(sql, params=()):
    with m.app.app_context():
        return m.get_db().execute(sql, params).fetchall()


def transcript_pane(html):
    part = html.split('id="pane-transcript"', 1)
    return part[1].split("<section", 1)[0] if len(part) > 1 else ""


admin = m.app.test_client()
check("landing 200", admin.get("/").status_code == 200)
check("anon /faculty redirects", admin.get("/faculty").status_code == 302)
check("anon /admin/faculty redirects", admin.get("/admin/faculty").status_code == 302)
r = login(admin, "admin", "Admin@123")
check("admin login", r.status_code == 302 and "/admin" in r.headers["Location"])
check("admin blocked from /faculty", admin.get("/faculty").status_code == 302)

# ------------------------------------------------------- fixtures (V7 namespace)
admin.post("/admin/departments/new", data={
    "name": "Dept of Faculty (V7)", "code": "FAC-V7", "faculty_name": "Faculty (V7)"})
admin.post("/admin/sessions/new", data={
    "name": "Spring 2027 (V7)", "starts_on": "2027-02-01", "ends_on": "2027-06-30",
    "is_open": "1"})
dept7 = dbq("SELECT id FROM departments WHERE code='FAC-V7'")[0]["id"]
sess7 = dbq("SELECT id FROM academic_sessions WHERE name='Spring 2027 (V7)'")[0]["id"]
admin.post("/admin/programs/new", data={
    "name": "BS Faculty (V7)", "department_id": dept7, "degree_level": "BS",
    "duration_semesters": 8})
prog7 = dbq("SELECT id FROM programs WHERE name='BS Faculty (V7)'")[0]["id"]
courses7 = {}
for code in ("V7-101", "V7-102"):
    admin.post("/admin/courses/new", data={
        "program_id": prog7, "semester_no": 1, "code": code,
        "title": f"Course {code} (V7)", "credit_theory": 3, "credit_lab": 0,
        "faculty_name": ""})
    courses7[code] = dbq("SELECT id FROM courses WHERE code=?", (code,))[0]["id"]
students7 = {}
for reg, name in [("V7-REG-001", "V7 Student One"), ("V7-REG-002", "V7 Student Two")]:
    admin.post("/admin/students/new", data={
        "registration_no": reg, "name": name, "department_id": dept7,
        "program_id": prog7, "session_id": sess7, "status": "Active",
        "current_semester": 1, "initial_password": "Test@1234"})
    students7[reg] = dbq("SELECT id FROM students WHERE registration_no=?", (reg,))[0]["id"]
s1, s2 = students7["V7-REG-001"], students7["V7-REG-002"]

# ------------------------------------------------------- faculty create + login
r = admin.post("/admin/faculty/new", data={
    "full_name": "V7 Teacher A", "email": "teacher.a@v7.example",
    "phone": "0300-1111111", "designation": "Lecturer", "department_id": dept7,
    "initial_password": "Teach@123"})
check("create faculty A redirects", r.status_code == 302)
lst = admin.get("/admin/faculty").get_data(as_text=True)
check("shown-once flash carries the initial password",
      "initial password: Teach@123" in lst and "shown ONCE" in lst)
fac_a = dbq("SELECT * FROM faculty WHERE email='teacher.a@v7.example'")
check("faculty A record exists", len(fac_a) == 1)
fac_a = fac_a[0]
usr_a = dbq("SELECT * FROM users WHERE faculty_id=?", (fac_a["id"],))
check("faculty A login auto-created (username=email, role=faculty)",
      len(usr_a) == 1 and usr_a[0]["username"] == "teacher.a@v7.example"
      and usr_a[0]["role"] == "faculty")

r = admin.post("/admin/faculty/new", data={
    "full_name": "Copy Cat", "email": "teacher.a@v7.example",
    "designation": "Lecturer", "department_id": dept7})
body = r.get_data(as_text=True)
check("duplicate email blocked",
      r.status_code == 200 and "already registered" in body
      and len(dbq("SELECT id FROM faculty WHERE email='teacher.a@v7.example'")) == 1)

r = admin.post("/admin/faculty/new", data={
    "full_name": "V7 Teacher B", "email": "teacher.b@v7.example",
    "designation": "Assistant Professor", "department_id": dept7})
check("create faculty B redirects (generated password)", r.status_code == 302)
lst = admin.get("/admin/faculty").get_data(as_text=True)
mm = re.search(r"initial password: (\S+) — shown ONCE", lst)
check("generated password shown once in flash", bool(mm))
pw_b = mm.group(1) if mm else ""
fac_b = dbq("SELECT * FROM faculty WHERE email='teacher.b@v7.example'")[0]

fa = m.app.test_client()
r = login(fa, "teacher.a@v7.example", "Teach@123")
check("faculty A login lands on /faculty",
      r.status_code == 302 and r.headers["Location"].endswith("/faculty"))
fb = m.app.test_client()
r = login(fb, "teacher.b@v7.example", pw_b)
check("faculty B login works with the generated shown-once password",
      r.status_code == 302 and r.headers["Location"].endswith("/faculty"))

# ------------------------------------------------------- assignments
def assign(fid, course_id, session_id=sess7):
    return admin.post(f"/admin/faculty/{fid}/assign",
                      data={"course_id": course_id, "session_id": session_id})


r = assign(fac_a["id"], courses7["V7-101"])
check("assign V7-101 to A", r.status_code == 302
      and len(dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
                  (fac_a["id"],))) == 1)
assign(fac_a["id"], courses7["V7-101"])
edit = admin.get(f"/admin/faculty/{fac_a['id']}/edit").get_data(as_text=True)
check("duplicate assignment (same faculty+course+session) blocked",
      "already assigned to V7 Teacher A" in edit
      and len(dbq("SELECT id FROM course_assignments WHERE faculty_id=? AND course_id=?",
                  (fac_a["id"], courses7["V7-101"]))) == 1)
assign(fac_b["id"], courses7["V7-101"])
edit_b = admin.get(f"/admin/faculty/{fac_b['id']}/edit").get_data(as_text=True)
check("second faculty blocked from the same course+session",
      "already assigned to V7 Teacher A" in edit_b
      and len(dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
                  (fac_b["id"],))) == 0)
assign(fac_b["id"], courses7["V7-102"])
check("B can take a different course in the same session",
      len(dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
              (fac_b["id"],))) == 1)
aid_a = dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
            (fac_a["id"],))[0]["id"]
aid_b = dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
            (fac_b["id"],))[0]["id"]

r = admin.post(f"/admin/faculty/{fac_a['id']}/delete")
lst = admin.get("/admin/faculty").get_data(as_text=True)
check("delete blocked while assignments exist",
      "still hold" in lst and "course assignment" in lst
      and len(dbq("SELECT id FROM faculty WHERE id=?", (fac_a["id"],))) == 1)

# ------------------------------------------------------- faculty dashboards
dash_a = fa.get("/faculty").get_data(as_text=True)
check("A dashboard lists ONLY own course",
      "V7-101" in dash_a and "V7-102" not in dash_a)
check("A dashboard shows enrolled count 2 + teacher identity",
      "2 student(s) enrolled" in dash_a and "V7 Teacher A" in dash_a
      and "Lecturer" in dash_a)
dash_b = fb.get("/faculty").get_data(as_text=True)
check("B dashboard lists ONLY own course",
      "V7-102" in dash_b and "V7-101" not in dash_b)

# ------------------------------------------------------- roster + Excel
roster = fa.get(f"/faculty/course/{aid_a}/students").get_data(as_text=True)
check("roster equals the enrolled cohort",
      "V7-REG-001" in roster and "V7-REG-002" in roster
      and "V7 Student One" in roster and "V7 Student Two" in roster
      and "CUVAS-DEMO" not in roster)
xl = fa.get(f"/faculty/course/{aid_a}/students.xlsx")
check("roster Excel downloads", xl.status_code == 200 and xl.data[:2] == b"PK")
from openpyxl import load_workbook
ws = load_workbook(io.BytesIO(xl.data)).active
xl_rows = list(ws.iter_rows(values_only=True))
check("roster Excel: header Reg No/Name/Signature + 2 rows",
      xl_rows[0] == ("Reg No", "Name", "Signature") and len(xl_rows) == 3
      and {r[0] for r in xl_rows[1:]} == {"V7-REG-001", "V7-REG-002"},
      str(xl_rows))

# ------------------------------------------------------- cross-faculty guards
check("B roster of A's assignment -> 404",
      fb.get(f"/faculty/course/{aid_a}/students").status_code == 404)
check("B roster xlsx of A's assignment -> 404",
      fb.get(f"/faculty/course/{aid_a}/students.xlsx").status_code == 404)
check("B marks GET of A's assignment -> 404",
      fb.get(f"/faculty/course/{aid_a}/marks").status_code == 404)
r = fb.post(f"/faculty/course/{aid_a}/marks", data={f"marks_{s1}": "55"})
check("B marks POST of A's assignment -> 404, nothing written",
      r.status_code == 404
      and dbq("SELECT id FROM results WHERE course_id=?", (courses7["V7-101"],)) == [])

# ------------------------------------------------------- faculty marks flow
marks_page = fa.get(f"/faculty/course/{aid_a}/marks").get_data(as_text=True)
check("faculty marks page has NO publish control",
      "Publish All" not in marks_page and "sheet/publish" not in marks_page)
r = fa.post("/admin/results/sheet/publish",
            query_string={"course_id": courses7["V7-101"], "session_id": sess7},
            data={"action": "publish"})
check("faculty POST to the admin publish route is role-blocked",
      r.status_code == 302)
r = fa.post(f"/faculty/course/{aid_a}/marks",
            data={f"marks_{s1}": "80", f"absent_{s2}": "on"})
check("faculty marks save redirects", r.status_code == 302)
row1 = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?",
           (s1, courses7["V7-101"]))
row2 = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?",
           (s2, courses7["V7-101"]))
check("faculty save writes shared results (80 -> A-/3.7, absent -> F)",
      len(row1) == 1 and row1[0]["marks_obtained"] == 80
      and row1[0]["grade"] == "A-" and row1[0]["grade_point"] == 3.7
      and len(row2) == 1 and row2[0]["is_absent"] == 1
      and row2[0]["grade"] == "F", str([dict(x) for x in row1 + row2]))
check("faculty-saved results stay UNPUBLISHED",
      row1[0]["is_published"] == 0 and row2[0]["is_published"] == 0)

stu1 = m.app.test_client()
login(stu1, "V7-REG-001", "Test@1234")
pane = transcript_pane(stu1.get("/student").get_data(as_text=True))
check("student does NOT see faculty marks while unpublished",
      "V7-101" not in pane)

admin.post("/admin/results/sheet/publish",
           query_string={"course_id": courses7["V7-101"], "session_id": sess7},
           data={"action": "publish"})
pane = transcript_pane(stu1.get("/student").get_data(as_text=True))
check("student sees the marks after ADMIN publishes",
      "V7-101" in pane and ">A-<" in pane, pane[:160])

marks_page = fa.get(f"/faculty/course/{aid_a}/marks").get_data(as_text=True)
check("published set: faculty marks page shows the lock note",
      "PUBLISHED" in marks_page and "locked" in marks_page)
r = fa.post(f"/faculty/course/{aid_a}/marks", data={f"marks_{s1}": "10"})
row1b = dbq("SELECT * FROM results WHERE student_id=? AND course_id=?",
            (s1, courses7["V7-101"]))[0]
check("faculty cannot edit after admin published (marks still 80)",
      row1b["marks_obtained"] == 80, str(dict(row1b)))

# ------------------------------------------------------- role guards
check("faculty blocked from /admin", fa.get("/admin").status_code == 302)
check("faculty blocked from /student", fa.get("/student").status_code == 302)
check("faculty blocked from /admin/faculty",
      fa.get("/admin/faculty").status_code == 302)
check("student blocked from /faculty", stu1.get("/faculty").status_code == 302)
check("student blocked from /admin/faculty",
      stu1.get("/admin/faculty").status_code == 302)

# ------------------------------------------------------- B: email sync, reset, deactivate
r = admin.post(f"/admin/faculty/{fac_b['id']}/edit", data={
    "full_name": "V7 Teacher B", "email": "teacher.b2@v7.example",
    "phone": "", "designation": "Assistant Professor",
    "department_id": dept7, "is_active": "1"})
check("edit B (email change) redirects", r.status_code == 302)
fb2 = m.app.test_client()
r = login(fb2, "teacher.b2@v7.example", pw_b)
check("login follows the edited email (username=email kept in sync)",
      r.status_code == 302 and r.headers["Location"].endswith("/faculty"))
fb_old = m.app.test_client()
check("old email no longer logs in",
      login(fb_old, "teacher.b@v7.example", pw_b).status_code == 200)

r = admin.post(f"/admin/faculty/{fac_b['id']}/reset-password")
check("reset-password redirects to the edit page", r.status_code == 302)
edit_b = admin.get(f"/admin/faculty/{fac_b['id']}/edit").get_data(as_text=True)
mm = re.search(r"New password: (\S+) — shown ONCE", edit_b)
check("reset password shown once on the edit page", bool(mm))
pw_b2 = mm.group(1) if mm else ""
fb3 = m.app.test_client()
r = login(fb3, "teacher.b2@v7.example", pw_b2)
check("login works with the reset password",
      r.status_code == 302 and r.headers["Location"].endswith("/faculty"))
fb_old2 = m.app.test_client()
check("pre-reset password rejected",
      login(fb_old2, "teacher.b2@v7.example", pw_b).status_code == 200)

r = admin.post(f"/admin/faculty/{fac_b['id']}/edit", data={
    "full_name": "V7 Teacher B", "email": "teacher.b2@v7.example",
    "phone": "", "designation": "Assistant Professor", "department_id": dept7})
check("deactivate B via edit (checkbox off)", r.status_code == 302)
fac_b_now = dbq("SELECT * FROM faculty WHERE id=?", (fac_b["id"],))[0]
fb4 = m.app.test_client()
r = login(fb4, "teacher.b2@v7.example", pw_b2)
check("deactivated faculty cannot log in",
      fac_b_now["is_active"] == 0 and r.status_code == 200
      and "Invalid username or password" in r.get_data(as_text=True))
admin.post(f"/admin/faculty/{fac_b['id']}/edit", data={
    "full_name": "V7 Teacher B", "email": "teacher.b2@v7.example",
    "phone": "", "designation": "Assistant Professor",
    "department_id": dept7, "is_active": "1"})
fb5 = m.app.test_client()
check("reactivated faculty can log in again",
      login(fb5, "teacher.b2@v7.example", pw_b2).status_code == 302)

# ------------------------------------------------------- unassign + delete
admin.post(f"/admin/faculty/{fac_b['id']}/assign/{aid_b}/remove")
check("unassign works", dbq("SELECT id FROM course_assignments WHERE faculty_id=?",
                            (fac_b["id"],)) == [])
r = admin.post(f"/admin/faculty/{fac_b['id']}/delete")
check("delete succeeds once assignments are gone",
      r.status_code == 302
      and dbq("SELECT id FROM faculty WHERE id=?", (fac_b["id"],)) == []
      and dbq("SELECT id FROM users WHERE faculty_id=?", (fac_b["id"],)) == [])
fb6 = m.app.test_client()
check("deleted faculty cannot log in",
      login(fb6, "teacher.b2@v7.example", pw_b2).status_code == 200)

# ------------------------------------------------------- demo seed intact
demo_fac = m.app.test_client()
r = login(demo_fac, "faculty", "Faculty@123")
check("demo teacher login (faculty/Faculty@123) works",
      r.status_code == 302 and r.headers["Location"].endswith("/faculty"))
dash = demo_fac.get("/faculty").get_data(as_text=True)
check("demo teacher has the 2 seeded assignments + identity",
      "Demo Teacher" in dash and "CS-201" in dash and "CS-202" in dash
      and "Lecturer" in dash)
demo_assign = dbq(
    """SELECT a.id FROM course_assignments a JOIN faculty f ON f.id=a.faculty_id
       WHERE f.email='demo.teacher@cuvas.edu.pk' ORDER BY a.id""")
check("demo teacher assignment rows == 2", len(demo_assign) == 2)
roster = demo_fac.get(
    f"/faculty/course/{demo_assign[0]['id']}/students").get_data(as_text=True)
check("demo roster is non-empty and includes the demo student",
      "CUVAS-DEMO-2026-001" in roster and "Demo Student" in roster)
demo_stu = m.app.test_client()
login(demo_stu, "student", "Student@123")
dh = demo_stu.get("/student").get_data(as_text=True)
check("demo student: 10 tabs + CGPA 3.63 untouched by v0.7",
      dh.count("data-tab=") == 10
      and "Overall CGPA: <strong>3.63</strong>" in transcript_pane(dh))

# ------------------------------------------- v0.6-era DB migration simulation
# Reconstruct a v0.6-era database (no faculty tables, no users.faculty_id,
# no faculty logins) and boot the v0.7 init against it, twice.
with m.app.app_context():
    raw = m.connect()
    raw.execute("DROP TABLE IF EXISTS course_assignments")
    raw.execute("DROP TABLE IF EXISTS faculty")
    if m.PG:
        raw.execute("ALTER TABLE users DROP COLUMN IF EXISTS faculty_id")
    else:
        raw.execute("ALTER TABLE users DROP COLUMN faculty_id")
    raw.execute("DELETE FROM users WHERE username='faculty'")
    raw.commit()
    raw.conn.close()
m.init_db()
fac = dbq("SELECT * FROM faculty")
usr = dbq("SELECT * FROM users WHERE username='faculty'")
check("migration: v0.6-era DB boots clean, demo teacher re-seeded + linked",
      len(fac) == 1 and fac[0]["email"] == "demo.teacher@cuvas.edu.pk"
      and len(usr) == 1 and usr[0].get("faculty_id") == fac[0]["id"])
n_fac = len(dbq("SELECT id FROM faculty"))
n_asg = len(dbq("SELECT id FROM course_assignments"))
m.init_db()
check("migration: second boot is a no-op",
      len(dbq("SELECT id FROM faculty")) == n_fac
      and len(dbq("SELECT id FROM course_assignments")) == n_asg)
demo_fac2 = m.app.test_client()
check("demo teacher login still works after migration sim",
      login(demo_fac2, "faculty", "Faculty@123").status_code == 302)

print()
if FAIL:
    print(f"{len(FAIL)} FAILURES: {FAIL}")
    sys.exit(1)
print("ALL CHECKS PASSED")
